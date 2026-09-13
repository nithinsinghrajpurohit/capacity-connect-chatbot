import io, os, re, json, base64

class DocumentAnalyzer:
    @staticmethod
    def extract_from_base64(file_data_b64, filename='uploaded_file'):
        clean_b64 = file_data_b64
        if 'data:' in file_data_b64 and ';base64,' in file_data_b64:
            header, clean_b64 = file_data_b64.split(';base64,')
        file_bytes = base64.b64decode(clean_b64)
        ext = os.path.splitext(filename)[1].lower() if filename else ''
        if ext == '.pdf' or b'%PDF' in file_bytes[:10]:
            return DocumentAnalyzer.extract_pdf(file_bytes)
        elif ext in ('.docx', '.doc'):
            return DocumentAnalyzer.extract_docx(file_bytes)
        else:
            try:
                text = file_bytes.decode('utf-8', errors='ignore')
                return {'type': 'text', 'filename': filename, 'text': text, 'pages': 1, 'char_count': len(text)}
            except Exception as e:
                return {'type': 'binary', 'filename': filename, 'text': f"Binary content ({len(file_bytes)} bytes)", 'pages': 1, 'char_count': 0}

    @staticmethod
    def extract_pdf(file_bytes):
        text_pages = []
        try:
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            for i, page in enumerate(reader.pages):
                txt = page.extract_text() or ''
                if txt.strip():
                    text_pages.append(f"--- Page {i + 1} ---\n{txt.strip()}")
        except Exception as e:
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
                    for i, page in enumerate(pdf.pages):
                        txt = page.extract_text() or ''
                        if txt.strip():
                            text_pages.append(f"--- Page {i + 1} ---\n{txt.strip()}")
            except Exception as e2:
                print(f"[DocumentAnalyzer] PDF extraction failed: {e2}")
        full_text = "\n\n".join(text_pages).strip()
        return {'type': 'pdf', 'pages': len(text_pages), 'text': full_text, 'char_count': len(full_text)}

    @staticmethod
    def extract_docx(file_bytes):
        paragraphs = []
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            for p in doc.paragraphs:
                if p.text and p.text.strip():
                    if p.style and 'heading' in p.style.name.lower():
                        paragraphs.append('## ' + p.text.strip())
                    else:
                        paragraphs.append(p.text.strip())
            for t_idx, table in enumerate(doc.tables):
                table_rows = []
                for row in table.rows:
                    row_cells = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    table_rows.append(' | '.join(row_cells))
                if table_rows:
                    paragraphs.append(f"\n[Table {t_idx + 1}]\n" + "\n".join(table_rows))
        except Exception as e:
            print(f"[DocumentAnalyzer] DOCX extraction failed: {e}")
        full_text = "\n\n".join(paragraphs).strip()
        return {'type': 'docx', 'pages': 1, 'text': full_text, 'char_count': len(full_text)}

    @staticmethod
    def build_analysis_prompt(doc_info, user_query='', user_name='Learner'):
        doc_type = doc_info.get('type', 'document').upper()
        doc_text = doc_info.get('text', '').strip()
        if len(doc_text) > 25000:
            doc_text = doc_text[:25000] + "\n\n[... Remaining content truncated for token limits ...]"
        pq = user_query or 'Provide a comprehensive breakdown, key takeaways, and curriculum insights.'
        return ("You are Sastra AI, performing a deep multimodal & technical document analysis for " + user_name + " on Capacity Connect (Government of Odisha).\n\n" +
                "The user uploaded a " + doc_type + " file for analysis.\n\n" +
                "User Question/Instructions: " + pq + "\n\n" +
                "--- DOCUMENT CONTENT ---\n" + doc_text + "\n--- END DOCUMENT CONTENT ---\n\n" +
                "Please provide a masterclass-tier document breakdown with the following sections:\n" +
                "1. 🐌 Executive Summary & Document Core Purpose\n" +
                "2. 🔔 Critical Insights & Conceptual Breakdown (Deep dive into all key sections, definitions, and logic)\n" +
                "3. 📩 Key Data Points, Formulas, Code, or Tables (if present)\n" +
                "4. 💡 Actionable Takeaways & Practical Relevance to the learner coursework\n" +
                "5. 🎯 Concept Check / Quiz Questions (3 short self-assessment questions based on the document)\n\n" +
                "Ensure zero raw asterisks in the response, formatting headers with clean symbols (✦, ◊, ♭, 💡).")

_analyzer = None
def get_document_analyzer():
    global _analyzer
    if _analyzer is None:
        _analyzer = DocumentAnalyzer()
    return _analyzer
